"""Report only observed V2 outcomes; never writes to V1."""
import csv
import json
import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path
import baseline_experiment as b
OUT=b.ROOT/'data/botou_yali/evaluation/baseline_v2'
OLD=b.ROOT/'data/botou_yali/evaluation/baseline_v1'


def avg(values):
    values=[v for v in values if v is not None]
    return statistics.mean(values) if values else None


def percent(x):return '—' if x is None else f'{100*x:.2f}%'


def calculate():
    manifest=b.read(OUT/'manifest.json')
    rows=[json.loads(line) for line in b.QA.read_text(encoding='utf-8').splitlines()]
    selected=[r for r in rows if r['id'] in manifest['selected_ids']]
    details=[]
    for row in selected:
        prepared=b.read(OUT/'retrieval'/(row['id']+'.json')) if (OUT/'retrieval'/(row['id']+'.json')).exists() else {}
        for arm in 'ABCD':
            path=OUT/'answers'/f"{row['id']}_{arm}.json"
            result=b.read(path) if path.exists() else {}
            judge=result.get('judge',{}).get('result')
            candidates=[] if arm=='A' else b.read(OUT/'vector'/(row['id']+'.json'))[:5] if arm=='B' else prepared.get(arm,[])[:5]
            metric=b.retrieval_metrics(row,candidates,5)
            details.append({'id':row['id'],'arm':arm,'question':row['question'],'type':row['type'],'source':row['source_group'],
                'reference_answer':row['reference_answer'],'evidence_kind':row['evidence_kind'],'gold_chunk':row['evidence']['chunk_id'],
                'url':row['evidence']['source_url'],'pages':row['evidence']['pages'],
                'finished':bool(result),'judged':bool(judge),'correct':int(bool(judge) and judge['correctness']=='correct'),
                'label':judge['correctness'] if judge else 'failed_or_missing','reason':judge['reason'] if judge else result.get('error','not_started'),
                'claims':judge.get('claims',[]) if judge else [],'answer':result.get('answer'),
                'citation_hit':result.get('citation_hit',0),'citation_source_hit':result.get('citation_source_hit',0),
                'entity_recall':result.get('entity_recall',0),'recall5':metric['recall'] if arm!='A' else None,
                'mrr5':metric['mrr'] if arm!='A' else None,'graph_top5':sum(c['kind']=='graph_path' for c in candidates),
                'status':result.get('service_response',{}).get('status'),'seconds':result.get('seconds')})
    arms={}
    for arm in 'ABCD':
        items=[r for r in details if r['arm']==arm];judged=[r for r in items if r['judged']]
        causal=[r for r in items if r['type']=='causal' and r['evidence_kind']=='reported_study_record']
        claims=Counter(c['status'] for r in judged for c in r['claims'])
        arms[arm]={'n':len(items),'finished':sum(r['finished'] for r in items),'judged':len(judged),
            'failed':sum(r['finished'] and not r['judged'] for r in items),'correct_n':sum(r['correct'] for r in items),
            'accuracy':avg([r['correct'] for r in items]),'recall5':avg([r['recall5'] for r in items]),'mrr5':avg([r['mrr5'] for r in items]),
            'citation_hit':avg([r['citation_hit'] for r in items]) if arm!='A' else None,
            'citation_source_hit':avg([r['citation_source_hit'] for r in items]) if arm!='A' else None,
            'entity_recall':avg([r['entity_recall'] for r in items]),'causal_n':len(causal),'causal_correct':sum(r['correct'] for r in causal),
            'causal_accuracy':avg([r['correct'] for r in causal]),
            'unsupported_rate_proxy':avg([int(any(c['status']!='supported' for c in r['claims'])) for r in judged]),
            'contradiction_rate':avg([int(any(c['status']=='contradicted' for c in r['claims'])) for r in judged]),
            'abstention_rate':avg([int(r['label']=='abstain') for r in judged]),'claim_counts':dict(claims),
            'graph_top5_questions':sum(r['graph_top5']>0 for r in items),'statuses':dict(Counter(r['status'] for r in items))}
    vector={}
    for k in (5,10,20):
        values=[b.retrieval_metrics(r,b.read(OUT/'vector'/(r['id']+'.json')),k) for r in rows]
        vector[k]={name:avg([m[name] for m in values]) for name in ('recall','mrr')}
    old=b.read(OLD/'summary.json')
    pairs={}
    keyed={(r['id'],r['arm']):r for r in details}
    for other in 'ABC':
        values=[(keyed[r['id'],'D']['correct'],keyed[r['id'],other]['correct']) for r in selected]
        pairs['D_vs_'+other]={'wins':sum(x>y for x,y in values),'ties':sum(x==y for x,y in values),'losses':sum(x<y for x,y in values)}
    unchanged=all((OLD/name).exists() and b.sha(OLD/name)==digest for name,digest in manifest['old_files'].items())
    summary={'time':datetime.now().isoformat(),'complete':all(r['finished'] for r in details) and (OUT/'completed.json').exists(),
             'old_files_unchanged':unchanged,'arms':arms,'old_accuracy':{a:old['arms'][a]['accuracy'] for a in 'ABCD'},
             'old_failures':{a:old['arms'][a]['failures'] for a in 'ABCD'},'vector350':vector,'paired':pairs,
             'type_counts':dict(Counter(r['type'] for r in selected)),
             'protocol_changed':True,'independent_test_set':False,'human_adjudicated':False}
    b.save(OUT/'summary.json',summary)
    (OUT/'per_question_metrics.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in details),encoding='utf-8')
    with (OUT/'human_review.csv').open('w',encoding='utf-8-sig',newline='') as f:
        fields=['id','arm','question','source','url','pages','gold_chunk','reference_answer','answer','label','reason','human_correctness','reviewer','notes']
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader()
        for row in details:writer.writerow({k:json.dumps(row[k],ensure_ascii=False) if k=='answer' else row.get(k,'') for k in fields})
    return summary,details


def report(summary,details):
    from docx import Document
    from docx.shared import Pt,Cm
    from docx.oxml.ns import qn
    doc=Document();style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.sections[0].left_margin=Cm(1.7);doc.sections[0].right_margin=Cm(1.7)
    def p(s):doc.add_paragraph(str(s))
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
        for row in rows:
            for c,s in zip(t.add_row().cells,row):c.text=str(s)
    doc.add_heading('农业 GraphRAG 新版 A/B/C/D 对比评估报告',0)
    p('V2独立目录重测 · 同一64题开发回归集 · '+summary['time'])
    p('实验完成：'+str(summary['complete'])+'；V1全部文件哈希保持不变：'+str(summary['old_files_unchanged']))
    p('本报告为模型辅助评分，不是人工专家验收。该64题集已用于诊断和修复，属于开发回归集，不再是独立盲测集。任何提升都不能直接作为未见问题上的泛化性能或比赛最终准确率。')
    doc.add_heading('1. 输入与实验设计',1)
    p('沿用V1固定64题号、3篇论文语料、同一BGE/FAISS索引及V1起始时读取的Neo4j真实快照。350题向量检索在V2重新运行；256份回答全部重新生成，未复用V1答案或评分。原始资料、题目和参考答案未修改。')
    p('64题包括37道factual题及全部27道少数类型题；gold答案与来源仅在答案完成后的评分环节使用。因果指标采用12道非背景causal题，排除1道背景/转引题；不证明SCM、干预效应或反事实能力。')
    table(['题型','题数'],summary['type_counts'].items())
    table(['组别','实现'],[
        ('A','DeepSeek闭卷，无检索、无外部证据。'),
        ('B','普通BGE+FAISS Vector Top20→按相似度取Top5→DeepSeek。'),
        ('C','新版实体入口与事实导航→无农业范围过滤的1～3跳路径+Vector Top20→BGE重排Top5→DeepSeek。非Microsoft GraphRAG官方复现。'),
        ('D','实际新版ask主链路：问题用途路由→实体对齐→双路检索→农业范围/证据准入→BGE+排名融合Top5→DeepSeek。')])
    p('B/C/D统一采用新版引用规范化、逐字段证据核查及最多一次反馈修订，避免只给D修复格式导致不公平。生成统一temperature=0、max_tokens=2500；独立评分Prompt沿用V1，max_tokens=2400。A无法对检索证据核查，保留闭卷设置。')
    p('图路径候选每组最多20条、向量最多20块，生成Top5；沿用V1候选预算，而线上默认图路径上限为100。D复用真实预计算检索与排名结果，调用生产ask函数处理生成、核查和输出；适配器只存在于独立实验进程，不修改在线服务。')
    p('新旧协议差异包括实体入口/导航修复、B/C共同引用核查、最大生成长度1600→2500、D用途路由与融合。新旧差异是系统版本回归，不可全部归因于农业约束算法本身。运行期间未调权重、换题或择优选择重复生成结果；失败结果同样保留。')
    doc.add_heading('2. 指标口径',1)
    p('正确率：模型盲评correct/64，partial、拒答与失败均不计正确。证据命中：最终正文引用的gold chunk命中题数/64，另列来源级命中。实体召回：答案逐字包含gold实体表面词的比例，失败计0，非图谱ID对齐准确率。')
    p('Recall@5与MRR@5对实际送入生成的Top5计分，相关项为已标注的gold chunk；每题仅一个gold块，未穷尽等价证据。来源相同的图路径不能冒充精确块命中。四组主表只比较K=5；不把D仅保留的Top5伪称Top20评测。')
    p('幻觉率代理：已评分回答中含至少一条contradicted或unverifiable断言的比例。unverifiable仅表示提供证据不足以验证，不等于客观虚假。另列明确矛盾率、拒答率和失败数。该代理不能脱离覆盖率解读。')
    p('评分模型与回答模型同源，制题也有模型参与，存在自评偏差。论文标题提示、题目近重复和仅三篇论文均限制外推；不声称统计显著提升。')
    doc.add_heading('3. 新版配对结果',1)
    arms=summary['arms']
    metrics=[('完成任务数','finished',False),('已评分数','judged',False),('失败数','failed',False),('正确题数','correct_n',False),
        ('端到端正确率','accuracy',True),('引用gold块命中率','citation_hit',True),('引用gold来源命中率','citation_source_hit',True),
        ('答案表面实体召回率','entity_recall',True),('Recall@5','recall5',True),('MRR@5','mrr5',False),
        ('幻觉率代理','unsupported_rate_proxy',True),('明确矛盾回答率','contradiction_rate',True),('拒答率（已评分）','abstention_rate',True),
        ('因果问题正确率（12题）','causal_accuracy',True),('Top5含图路径题数','graph_top5_questions',False)]
    table(['指标','A','B','C','D'],[[label]+[percent(arms[a][key]) if ratio else ('—' if arms[a][key] is None else round(arms[a][key],4)) for a in 'ABCD'] for label,key,ratio in metrics])
    doc.add_heading('4. 与V1的版本回归比较',1)
    table(['组别','V1正确率','V2正确率','变化/百分点','V1失败','V2失败'],[
        [a,percent(summary['old_accuracy'][a]),percent(arms[a]['accuracy']),f"{100*(arms[a]['accuracy']-summary['old_accuracy'][a]):+.2f}",summary['old_failures'][a],arms[a]['failed']] for a in 'ABCD'])
    for key,value in summary['paired'].items():p(f"新版{key}：D胜{value['wins']}题、相同{value['ties']}题、负{value['losses']}题。")
    p('这些差异包含共同的工程修复、证据核查及模型生成波动。只有另做单模块消融和独立专家测试，才能分别论证检索、约束与核查的贡献。')
    doc.add_heading('5. 全量350题向量检索',1)
    table(['K','Recall@K','MRR@K'],[[k,percent(v['recall']),round(v['mrr'],4)] for k,v in summary['vector350'].items()])
    p('只有检索覆盖350题，四组完整问答仍为相同64题。不得宣称350题四组全部生成评测。')
    doc.add_heading('6. 失败、图谱参与和典型差异',1)
    for a in 'CD':p(f"{a}的Top5包含图路径：{arms[a]['graph_top5_questions']}/64。若D为0，论文阅读效果不能解释为图谱因果推理有效；标题限定可能正确排除非论文来源的路径。")
    p('失败与缺失任务：')
    for row in details:
        if not row['judged']:p(f"{row['id']} {row['arm']}：{row['reason']}")
    p('按题号展示前6个D/B正确性不同的案例，不按有利方向选取；全部结果见逐题文件。')
    ids=sorted({r['id'] for r in details});count=0
    for qid in ids:
        group={r['arm']:r for r in details if r['id']==qid}
        if group['D']['correct']==group['B']['correct']:continue
        p(qid+' '+group['D']['question']);p('参考答案：'+group['D']['reference_answer'])
        for a in 'BD':p(a+'：'+group[a]['label']+'；'+group[a]['reason'])
        count+=1
        if count==6:break
    doc.add_heading('7. 复现与后续验证',1)
    p('独立目录：'+str(OUT.resolve()))
    p('运行：ai-service/.venv/Scripts/python.exe -u ai-service/scripts/baseline_v2_experiment.py')
    p('报告：ai-service/.venv/Scripts/python.exe ai-service/scripts/baseline_v2_report.py')
    p('manifest冻结题号、语料/图谱/代码哈希及旧目录全部文件哈希；code_snapshot保留评测时实现。已完成结果包含失败，恢复运行时不会反复抽取更好的答案。API密钥不进入记录。')
    p('answers保留发布回答、独立评分、生成请求标识和修订；retrieval保留候选与剔除原因；human_review.csv包含原文位置与人工审核空栏。生成与评分返回的token记录不含完整核查/抽取用量，不能冒充总账单。')
    p('下一步需要独立专家审定的地域/作物冲突、证据不足、多跳及生产—采后问题集。当前64题只能作为开发回归，不作为未见数据的最终性能证明。')
    p('方法沿用V1：RAGAS https://arxiv.org/abs/2309.15217 与BEIR指标 https://github.com/beir-cellar/beir/wiki/Metrics-available；采用自定义、有参考答案的评审，不冒称官方RAGAS分数。')
    path=OUT/'新版农业GraphRAG_ABCD对比评估报告.docx';doc.save(path)
    print('REPORT',str(path.resolve()))
    print(json.dumps({'complete':summary['complete'],'accuracy':{a:v['accuracy'] for a,v in arms.items()},'V1_unchanged':summary['old_files_unchanged']},ensure_ascii=False))


if __name__=='__main__':report(*calculate())
