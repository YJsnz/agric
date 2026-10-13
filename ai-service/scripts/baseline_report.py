"""Compute metrics from persisted observations; emit audit files and a Chinese DOCX."""
import csv
import json
import statistics
import importlib.metadata
import platform
from collections import Counter
from datetime import datetime
from pathlib import Path
from baseline_experiment import OUT, QA, ROOT, read, save, retrieval_metrics, sha


def mean(values):
    values = [v for v in values if v is not None]
    return statistics.mean(values) if values else None


def pct(value):
    return '不适用/无结果' if value is None else f'{100 * value:.2f}%'


def calculate():
    manifest = read(OUT / 'manifest.json')
    rows = [json.loads(s) for s in QA.read_text(encoding='utf-8').splitlines()]
    chosen = [r for r in rows if r['id'] in manifest['selected_ids']]
    details = []
    for row in chosen:
        path = OUT / 'retrieval' / (row['id'] + '.json')
        retrieval = read(path) if path.exists() else {}
        for arm in 'ABCD':
            file = OUT / 'answers' / f"{row['id']}_{arm}.json"
            answer = read(file) if file.exists() else {}
            judge = answer.get('judge', {}).get('result')
            candidates = [] if arm == 'A' else read(OUT / 'vector' / (row['id'] + '.json')) if arm == 'B' else retrieval.get(arm, [])
            result = {'id': row['id'], 'arm': arm, 'type': row['type'], 'source': row['source_group'],
                      'evidence_kind': row['evidence_kind'], 'question': row['question'],
                      'gold_chunk': row['evidence']['chunk_id'], 'reference_answer': row['reference_answer'],
                      'paper_title': row['evidence']['title'], 'gold_source_url': row['evidence']['source_url'],
                      'gold_pages': row['evidence']['pages'],
                      'completed': bool(judge), 'error': answer.get('error') or (retrieval.get('error') if arm in 'CD' else None),
                      'correct': int(judge['correctness'] == 'correct') if judge else 0,
                      'judge_label': judge['correctness'] if judge else 'missing',
                      'judge_reason': judge['reason'] if judge else '',
                      'citation_hit': answer.get('citation_gold_chunk_hit', 0),
                      'citation_source_hit': answer.get('citation_gold_source_hit', 0),
                      'entity_recall': answer.get('answer_entity_surface_recall', 0),
                      'answer': answer.get('generation', {}).get('result'),
                      'claims': judge.get('claims', []) if judge else [],
                      'generation_seconds': answer.get('generation', {}).get('seconds'),
                      'retrieval_seconds_shared': retrieval.get('seconds') if arm in 'CD' else None,
                      'graph_top5_count': sum(c['kind']=='graph_path' for c in candidates[:5])}
            result['attempt_finished'] = bool(judge or result['error'])
            for k in (5, 10, 20):
                m = retrieval_metrics(row, candidates, k)
                for name, value in m.items():
                    result[f'{name}@{k}'] = value if arm != 'A' else None
            details.append(result)
    arms = {}
    for arm in 'ABCD':
        entries = [r for r in details if r['arm'] == arm]
        judged = [r for r in entries if r['completed']]
        causal = [r for r in entries if r['type']=='causal' and r['evidence_kind']=='reported_study_record']
        claims = [c for r in judged for c in r['claims']]
        counts = Counter(c['status'] for c in claims)
        arms[arm] = {'n': len(entries), 'judged_n': len(judged), 'failures': len(entries)-len(judged),
            'correct_n': sum(r['correct'] for r in entries), 'accuracy': mean([r['correct'] for r in entries]),
            'correctness_labels': dict(Counter(r['judge_label'] for r in entries)),
            'answer_entity_surface_recall': mean([r['entity_recall'] for r in entries]),
            'evidence_chunk_citation_hit': mean([r['citation_hit'] for r in entries]) if arm != 'A' else None,
            'evidence_source_citation_hit': mean([r['citation_source_hit'] for r in entries]) if arm != 'A' else None,
            'causal_n': len(causal), 'causal_correct_n': sum(r['correct'] for r in causal),
            'causal_accuracy': mean([r['correct'] for r in causal]),
            'unsupported_answer_rate_proxy': mean([int(any(c['status'] != 'supported' for c in r['claims'])) for r in judged]),
            'contradiction_answer_rate': mean([int(any(c['status'] == 'contradicted' for c in r['claims'])) for r in judged]),
            'claim_counts': dict(counts), 'unsupported_claim_rate_proxy': (counts['contradicted']+counts['unverifiable'])/len(claims) if claims else None,
            'abstention_rate': mean([int(r['judge_label']=='abstain') for r in judged]),
            'graph_top5_questions': sum(r['graph_top5_count'] > 0 for r in entries),
            'by_type': {t: {'n': sum(r['type']==t for r in entries),
                           'accuracy': mean([r['correct'] for r in entries if r['type']==t])} for t in sorted({r['type'] for r in entries})}}
        for k in (5, 10, 20):
            for name in ('recall','mrr'):
                arms[arm][f'{name}@{k}'] = mean([r[f'{name}@{k}'] for r in entries])
    all_vector = {}
    for k in (5, 10, 20):
        values = [retrieval_metrics(r, read(OUT/'vector'/(r['id']+'.json')), k) for r in rows]
        all_vector[k] = {name: mean([v[name] for v in values]) for name in ('recall','mrr')}
    comparisons = {}
    by_id = {(r['id'], r['arm']): r for r in details}
    for other in 'ABC':
        pairs = [(by_id[r['id'],'D']['correct'], by_id[r['id'],other]['correct']) for r in chosen]
        comparisons['D_vs_'+other] = {'wins':sum(a>b for a,b in pairs),'ties':sum(a==b for a,b in pairs),
                                     'losses':sum(a<b for a,b in pairs),'accuracy_difference':mean([a-b for a,b in pairs])}
    usage = Counter()
    for file in (OUT/'answers').glob('*.json'):
        result = read(file)
        for kind in ('generation','judge'):
            for key in ('prompt_tokens','completion_tokens','total_tokens'):
                usage[key] += result.get(kind,{}).get('usage',{}).get(key,0)
    summary = {'generated_at':datetime.now().isoformat(), 'manifest':manifest, 'arms':arms,
               'vector_all_350':all_vector,'paired_comparisons':comparisons,'recorded_api_usage':dict(usage),
               'usage_note':'仅成功生成/评审请求返回的用量；不含实体抽取、失败重试，不能当账单总量',
               'complete':len(details)==256 and all(r['attempt_finished'] for r in details),
               'all_answers_judged': all(r['completed'] for r in details)}
    summary['indexed_corpus'] = {k:v for k,v in read(ROOT/'data/botou_yali/vector/metadata.json').items()
                                 if k in ('model','chunk_count','vector_count','dimension','model_sha256','tokenizer_sha256')}
    summary['selected_evidence_clusters'] = len({r['evidence']['chunk_id'] for r in chosen})
    summary['models_observed'] = dict(Counter(read(f).get('generation',{}).get('model')
                                            for f in (OUT/'answers').glob('*.json')))
    retrieval_rows = [read(f) for f in (OUT/'retrieval').glob('*.json')]
    summary['retrieval_diagnostics'] = {
        'completed_questions':len(retrieval_rows),
        'failed_questions':sum('error' in r for r in retrieval_rows),
        'C_graph_candidate_questions':sum(r.get('C_graph_count',0)>0 for r in retrieval_rows),
        'D_graph_candidate_questions':sum(r.get('D_graph_count',0)>0 for r in retrieval_rows),
        'D_empty_questions':sum('error' not in r and not r.get('D') for r in retrieval_rows),
        'document_exclusion_occurrences':dict(Counter(reason for r in retrieval_rows
            for e in r.get('excluded_documents',[]) for reason in e['reasons']))}
    provenance = {'python':platform.python_version(), 'platform':platform.platform(),
                  'packages':{name:importlib.metadata.version(name) for name in ('python-docx','fastembed','onnxruntime','faiss-cpu','numpy')},
                  'source_hashes_at_report':{name:sha(ROOT/'scripts'/name) for name in ('baseline_experiment.py','baseline_report.py','test_baseline_experiment.py')},
                  'reranker_manifest':read(ROOT/'models/reranker/manifest.json')}
    save(OUT/'runtime_provenance.json', provenance)
    save(OUT/'summary.json',summary)
    (OUT/'per_question_metrics.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in details),encoding='utf-8')
    with (OUT/'human_review.csv').open('w',encoding='utf-8-sig',newline='') as f:
        fields=['id','arm','question','paper_title','gold_source_url','gold_pages','gold_chunk','reference_answer','answer','judge_label','judge_reason','human_correctness','human_hallucination','reviewer','notes']
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader()
        for r in details:
            writer.writerow({k:json.dumps(r[k],ensure_ascii=False) if k=='answer' else r.get(k,'') for k in fields})
    return summary, details


def report(summary, details):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    section=doc.sections[0];section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    footer=section.footer.paragraphs[0];footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def p(text):doc.add_paragraph(str(text))
    def table(headers, rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        return t
    doc.add_heading('农业 GraphRAG Baseline 对比评估报告',0)
    p('第十五阶段 · 实测初步报告 · 350题向量检索 / 64题四组对比')
    p('生成时间：'+summary['generated_at'])
    p('状态：'+('256个实验任务均已结束，失败项计入统计' if summary['complete'] else '实验尚未全部结束，不视为完整实验'))
    p('通过校验并完成模型辅助评分：'+str(sum(v['judged_n'] for v in summary['arms'].values()))+'/256。接口失败与评分缺失详见统计；不得将失败题从正确率分母中删除。')
    p('本报告不预设 D 优于 A/B/C。答案正确率和幻觉代理由同一 DeepSeek 模型辅助评审，尚无独立专家审定；不得写成正式人工准确率或比赛最终性能。')
    doc.add_heading('1. 数据与实验边界',1)
    p('原题库350题来自PAPER_001、PAPER_003、PAPER_005三篇论文。333题为研究记录，17题为背景或转引讨论。参考答案可逐字定位原始块，但这不等于实验结论本身已被独立验证。题目含论文标题，来源检索较容易，不能代表真实农户问题。')
    p('预先固定64题：全部27道非factual题，加37道按SHA-256排序、三篇论文轮转抽取的factual题。选择种子20260921，未依据任何运行结果选题。未训练或调权重；因此本次为固定评测集初步测量，不是训练/验证/测试随机切分。')
    p('同一论文与证据块的题目相关，三篇论文不足以可靠做跨地区泛化推断。本报告不声称统计显著提升，不使用按题独立假设的置信区间。')
    p('64题仅覆盖'+str(summary['selected_evidence_clusters'])+'个不同gold证据块；不能当作64个独立农业实验。')
    table(['类型','64题数量'],summary['manifest']['type_counts'].items())
    p('因果标签共13题，其中1题为背景/转引，排除后以12题计分。该指标仅表示对论文影响/机制问题的回答正确性，不证明结构因果模型、do干预或反事实识别能力。环境风险与干预建议题型无可用专属题目。')
    doc.add_heading('2. 四种方案与控制变量',1)
    table(['组别','实现'],[
        ('A','DeepSeek闭卷回答；无文档、无图谱、无外部检索。允许承认不知道论文内容。'),
        ('B','BGE Embedding + FAISS Top20，按余弦相似度取Top5，再由DeepSeek回答；无重排与农业约束。'),
        ('C','本项目普通图检索对照：实体抽取/对齐 + 无地域作物阶段过滤的1～3跳知识路径 + Vector Top20 + BGE重排Top5。不是Microsoft GraphRAG官方实现。'),
        ('D','农业约束GraphRAG：共享实体抽取 + 地域/作物/阶段约束子图与向量检索 + A/B/C证据等级、来源和审阅状态准入 + BGE重排Top5。')])
    p('C/D均限定最多20条图路径、20个向量候选，输出Top5。生产接口的图路径上限更大，因此这是明确预算的实验变体，不等同于完整在线接口延迟评测。C仍要求图谱边可追溯，未人为注入错误资料。')
    p('C/D采用同一BGE模型、350字符窗口及50字符重叠、窗口最大分聚合，共享同一候选的评分以节约CPU；D只有准入后的候选参与排名。两组允许检索不同子图，属于约束模块整体消融，不是单一权重消融。')
    p('B/C/D共享生产Prompt的证据约束、五段回答字段和引用校验。A改为闭卷指令，其他输出字段一致。所有生成temperature=0、max_tokens=1600；评审max_tokens=2400。API实际模型名与响应ID保存在逐条记录。')
    p('成功生成记录中的实际模型标识及次数（null为无模型调用的拒答）：'+json.dumps(summary['models_observed'],ensure_ascii=False))
    p('参考答案、gold来源、gold实体及评审原文只用于评分，未输入检索或回答生成。Neo4j在实验起始时读取并冻结快照，不使用前端模拟图。向量索引不包含本次测试题与答案。')
    corpus=summary['indexed_corpus']
    p(f"索引：{corpus['model']}，{corpus['dimension']}维，{corpus['chunk_count']}个原始块、{corpus['vector_count']}个窗口向量。重排：BAAI/bge-reranker-base，本地CPU ONNX。环境版本、模型下载清单及脚本当前哈希见runtime_provenance.json。")
    doc.add_heading('3. 指标定义与分母',1)
    table(['指标','定义与限制'],[
        ('答案正确率','模型盲评correct题数/64。partial不计正确，缺失与失败计0；拒答在可回答题中计0。'),
        ('证据命中率','回答实际引用的E编号映射到gold chunk的题数/64，失败计0；同时报告较宽松的gold source命中率。A没有检索引用，记不适用。'),
        ('实体召回率','回答中逐字出现的去重gold实体数/gold实体数，再按题平均。gold实体尚未标准化，只能称答案表面实体召回，不等于图谱实体对齐正确率。'),
        ('Recall@K','TopK中命中的已标注gold chunk数/gold chunk数，按题平均。每题目前只有1个gold块，因此等价于已标注块Hit@K；其他合理块未补标。'),
        ('MRR@K','前K个候选中第一个gold块排名的倒数，未命中为0，再平均；图路径占候选位置，但仅来源相同不算精确块命中。'),
        ('幻觉率代理','至少包含1条证据不支持断言的回答/已评审回答。证据不支持=contradicted或unverifiable，并非均已证实为虚假。另列明确矛盾率与拒答率。'),
        ('因果问题正确率','12道非背景的causal标签题中correct数量/12；仅为论文问答能力代理。')])
    p('评审输入为匿名单份回答、gold原文块和该回答实际检索证据，不提供A/B/C/D身份。评审仍使用同模型，存在自评偏差；题库也由该模型辅助整理。因此自动分数仅供定位问题，必须配合人工复核。')
    doc.add_heading('4. 64题配对实验结果',1)
    arms=summary['arms']
    table(['指标','A','B','C','D'],[
        ['完成评审数']+[f"{arms[a]['judged_n']}/64" for a in 'ABCD'],
        ['正确题数']+[f"{arms[a]['correct_n']}/64" for a in 'ABCD'],
        ['答案正确率（模型辅助）']+[pct(arms[a]['accuracy']) for a in 'ABCD'],
        ['引用gold块命中率']+[pct(arms[a]['evidence_chunk_citation_hit']) for a in 'ABCD'],
        ['引用gold来源命中率']+[pct(arms[a]['evidence_source_citation_hit']) for a in 'ABCD'],
        ['答案表面实体召回率']+[pct(arms[a]['answer_entity_surface_recall']) for a in 'ABCD'],
        ['Recall@5']+[pct(arms[a]['recall@5']) for a in 'ABCD'],
        ['Recall@10']+[pct(arms[a]['recall@10']) for a in 'ABCD'],
        ['Recall@20']+[pct(arms[a]['recall@20']) for a in 'ABCD'],
        ['MRR@5']+['—' if arms[a]['mrr@5'] is None else f"{arms[a]['mrr@5']:.4f}" for a in 'ABCD'],
        ['幻觉率代理（证据不支持）']+[pct(arms[a]['unsupported_answer_rate_proxy']) for a in 'ABCD'],
        ['明确证据矛盾回答率']+[pct(arms[a]['contradiction_answer_rate']) for a in 'ABCD'],
        ['拒答率']+[pct(arms[a]['abstention_rate']) for a in 'ABCD'],
        ['因果正确题数']+[f"{arms[a]['causal_correct_n']}/{arms[a]['causal_n']}" for a in 'ABCD'],
        ['因果问题正确率（代理）']+[pct(arms[a]['causal_accuracy']) for a in 'ABCD'],
        ['Top5含图路径的题数']+[str(arms[a]['graph_top5_questions']) if a in 'CD' else '—' for a in 'ABCD']])
    p('B/C/D的Recall@10/@20来自候选排名，回答仅使用Top5。不同K不能混为同一生成上下文。失败项不会伪造成正常的证据不足回答。')
    p('幻觉代理只以成功评审的回答为分母，必须结合完成率和拒答率解读；失败输出被接口拦截不等于模型不会产生错误。实体召回率中无输出的失败题计0，表面匹配不接受同义词替换。')
    table(['题型','n','A正确率','B正确率','C正确率','D正确率'],[
        [t,arms['A']['by_type'][t]['n']]+[pct(arms[a]['by_type'][t]['accuracy']) for a in 'ABCD']
        for t in arms['A']['by_type']])
    p('分题型表保留原始标签，causal含1道背景题（共13题）；主表因果指标已排除此题，仅12题。')
    doc.add_heading('5. 全部350题的Vector检索结果',1)
    table(['K','已标注块Recall@K','MRR@K'],[[k,pct(v['recall']),f"{v['mrr']:.4f}"] for k,v in summary['vector_all_350'].items()])
    p('此表仅为全部350题的向量检索，不是350题A/B/C/D完整生成实验。64题少数题型过采样，不能把四组总体准确率直接当350题总体准确率。')
    doc.add_heading('6. 差异、失败与结论',1)
    for arm in 'CD':
        p(f"{arm}组Top5实际包含图路径：{arms[arm]['graph_top5_questions']}/64题。若为0，不能把本组成绩解释为图谱推理有效；C此时实质退化为Vector+BGE重排。")
    diagnostics=summary['retrieval_diagnostics']
    p(f"D组通过约束后无候选：{diagnostics['D_empty_questions']}题；图谱检索阶段失败：{diagnostics['failed_questions']}题。")
    table(['文档排除原因','候选出现次数（非题数）'],diagnostics['document_exclusion_occurrences'].items())
    if (OUT/'graph_canary.json').exists():
        canary=read(OUT/'graph_canary.json')
        p(f"另作不计分功能检查：‘黑星病’对齐到{canary['seed_count']}个种子，检索到{canary['edge_count']}条边、{canary['path_count']}条路径。该检查证明检索器能工作，但不作为补加测试题，也不能替代本题集中的图谱参与率。")
    for name,c in summary['paired_comparisons'].items():
        p(f"{name}：D多答对{c['wins']}题，少答对{c['losses']}题，相同{c['ties']}题；正确率差{c['accuracy_difference']*100:+.2f}个百分点。")
    p('这些是当前固定阅读题集上的描述性差异。若D落后，不得删题或调参后沿用本测试集宣称独立提升；需检查适用阶段/地区元数据是否缺失、约束是否把论文阅读问题误当成本地生产建议、图路径是否挤占有效文档位置。')
    p('即使D更高，也不能从本题集推出地域幻觉抑制能力：现有题目不是成对的跨地区/跨品种冲突测试。应另建专家审定的范围冲突、证据不足及多跳问题集，按论文/实验簇划分验证与测试后再调权重。')
    failures=[r for r in details if not r['completed']]
    p('缺失或失败记录数：'+str(len(failures)))
    citation_failures=sum('ExtractionError,ExtractionError,ExtractionError' in (r['error'] or '') for r in failures)
    extraction_failures=sum('entity_extraction_failed' in (r['error'] or '') for r in failures)
    p(f'其中回答引用校验重试后仍失败{citation_failures}项；实体抽取失败导致下游未执行{extraction_failures}项。诊断样本出现正文写E1而非[E1]、正文与citations列表不一致等格式问题，属于接口契约失败，不能直接等同为农学知识错误。当前表中的正确率衡量端到端有效回答。')
    for r in failures:p(f"{r['id']} {r['arm']}：{r['error']}")
    p('优先修复方向：验证引用编号结构化输出与校验器的兼容性；人工核查采后/储藏/贮藏/货架等阶段的实体归一与元数据；区分论文阅读与当地生产建议的适用性策略。修复后使用新实验版本保存结果，不能覆盖此次失败记录来制造提升。')
    p('为避免只报有利案例，下列展示按题号排序的前6个C/D正确性不一致案例。完整256条评分、回答与人工复核空表一并交付。')
    seen=0
    for qid in summary['manifest']['selected_ids']:
        group={r['arm']:r for r in details if r['id']==qid}
        if group['C']['correct']==group['D']['correct']:continue
        p(qid+' '+group['C']['question'])
        p('参考：'+group['C']['reference_answer'])
        for arm in 'CD':p(arm+'：'+group[arm]['judge_label']+'；'+group[arm]['judge_reason'])
        seen+=1
        if seen>=6:break
    doc.add_heading('7. 复现与交付文件',1)
    p('实验目录：'+str(OUT.resolve()))
    p('运行：ai-service/.venv/Scripts/python.exe ai-service/scripts/baseline_experiment.py --workers 4')
    p('生成报告：ai-service/.venv/Scripts/python.exe ai-service/scripts/baseline_report.py')
    p('已有结果断点复用；manifest固定题号、语料与索引哈希，变化时拒绝混跑。修改实验代码或Prompt后必须使用新实验目录，不能混用旧缓存。')
    for name in ['manifest.json','graph_snapshot.json','summary.json','per_question_metrics.jsonl','human_review.csv']:
        p(name+'：'+sha(OUT/name))
    p('成功请求记录的token用量：'+json.dumps(summary['recorded_api_usage'],ensure_ascii=False)+'。'+summary['usage_note'])
    p('answers/保留每题回答、证据、API响应标识及评审理由；retrieval/保留排名、剔除原因及实体对齐；vector/保存350题向量结果。密钥不写入报告和结果。')
    doc.add_heading('8. 方法参考',1)
    p('Es et al., Ragas: Automated Evaluation of Retrieval Augmented Generation. https://arxiv.org/abs/2309.15217 。参考其分开评估检索、回答与忠实性的思路；本实验使用自定义、有参考答案的盲评Prompt，不冒称官方Ragas分数。')
    p('BEIR Metrics available: https://github.com/beir-cellar/beir/wiki/Metrics-available 。采用Recall@K和MRR@K定义；相关性标注仅覆盖当前gold块。')
    target=OUT/'农业GraphRAG_Baseline评估报告.docx'
    doc.save(target)
    print('REPORT',str(target.resolve()))
    print(json.dumps({'complete':summary['complete'],'accuracy':{a:v['accuracy'] for a,v in arms.items()}},ensure_ascii=False))


if __name__=='__main__':
    report(*calculate())
