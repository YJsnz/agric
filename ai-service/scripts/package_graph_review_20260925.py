"""Package source snapshots, extraction and a bounded review report."""
import json,hashlib,sys,urllib.request
from pathlib import Path
from html.parser import HTMLParser
from collections import Counter
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data/botou_yali/evidence/review_20260925'
class Text(HTMLParser):
    def __init__(self):super().__init__();self.parts=[];self.skip=0
    def handle_starttag(self,t,a):
        if t in ('script','style'):self.skip+=1
        if t in ('p','h1','h2','h3','h4','section'):self.parts.append('\n')
    def handle_endtag(self,t):
        if t in ('script','style'):self.skip=max(0,self.skip-1)
    def handle_data(self,d):
        if not self.skip:self.parts.append(d)
urls={'yali2024':'https://pmc.ncbi.nlm.nih.gov/articles/PMC11011782/','yali2025':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12191683/','moa':'https://scs.moa.gov.cn/ccll/jszn/202302/t20230203_6419771.htm','ucscab':'https://ipm.ucanr.edu/agriculture/pear/pear-scab/'}
manifest=[]
for key,url in urls.items():
    path=OUT/(key+'.html')
    if not path.exists():path=OUT/(key+'.download-blocked.html')
    raw=path.read_bytes();encoding='gb18030' if key=='moa' else 'utf-8'
    text=Text();text.feed(raw.decode(encoding,errors='replace'));body=' '.join(text.parts)
    blocked=key=='yali2025' and '2.1. Plant Material' not in body
    target=OUT/(key+('.download-blocked.html' if blocked else '.html'))
    if blocked:
        target.write_bytes(raw)
        if path!=target:path.unlink()
    else:(OUT/(key+'.txt')).write_text(body,encoding='utf-8')
    manifest.append(dict(source_id=key,url=url,accessed_at='2026-09-25',file=target.name,sha256=hashlib.sha256(raw).hexdigest(),local_fulltext=not blocked,status='browser_fulltext_read_local_download_blocked' if blocked else 'downloaded_body_verified',text_file='' if blocked else key+'.txt'))
(OUT/'sources.jsonl').write_text('\n'.join(json.dumps(s,ensure_ascii=False) for s in manifest)+'\n',encoding='utf-8')
g=json.loads((ROOT/'knowledge_graph/data/graph-v2.json').read_text(encoding='utf-8'))
before=json.loads((OUT/'graph-before.json').read_text(encoding='utf-8'))
with urllib.request.urlopen('http://localhost:8090/api/graph',timeout=30) as r:live=json.load(r)
audit={'review_date':'2026-09-25','scope':'结构全量检查；四份原文重点核对，不是307条旧关系逐条专家复审','before':{k:len(before[k]) for k in ('nodes','edges','sources')},'after':{k:len(g[k]) for k in ('nodes','edges','sources')},'live_matches_disk':all({v['id'] for v in live[k]}=={v['id'] for v in g[k]} for k in ('nodes','edges','sources')),'edge_review_status':dict(Counter(e['review_status'] for e in g['edges'])),'missing_evidence':[e['id'] for e in g['edges'] if not e.get('evidence_ids')],'unverified_causal_probability':[e['id'] for e in g['edges'] if e.get('confidence') is not None or e.get('effect_size') is not None]}
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
paragraphs=[
('结论与核对边界','不能宣称图谱全部内容已经科学验证。本次对全量数据做结构检查，对农业农村部、UC IPM及两篇鸭梨论文做重点证据核对。旧关系仍需农业专家逐条复核；来源可靠不等于因果效应成立。'),
('前端发现','泊头鸭梨页面包含一般梨类和异地资料。分类导航虚线不是数据库农业事实边。已增加待本地适用性审核提示，防止把分类归属误读成本地产生的结论。显示关系总数包含溯源，实际农业语义边315条。'),
('黑星病范围','UC IPM描述的是加州梨及Venturia pirina。保留异地适用性标记；不能把它无条件当成泊头鸭梨病原鉴定或当地用药依据。黑星病到采后腐烂的普遍因果链未获本次材料支持，不补造。'),
('采后规范','农业农村部区分侵染性病害和冷害、二氧化碳伤害等生理问题，并单列鸭梨温控要求。不同品种气体指标不能移植；本文未新增农药或消毒剂剂量建议。'),
('新增论文知识','新增2份独立论文来源、9个实体、8条关系。2024研究支持特定高成熟度鸭梨快速降温组的若干品质观察；2025研究比较不同成熟度在缓冷条件下的表现。均来自辛集试材，标记transfer_pending，不输出泊头风险概率。方法条件逐条附在conditions中。'),
('科学表达限制','2025文中部分描述存在表述不一致风险，本次仅抽取方法及结果3.1的限定比较，不采用其前言中的泛化机制，不计算因果效应。两文的成熟度和降温方案不同，不构成简单互相否定，也不能据此推荐所有鸭梨快冷。'),
('文档存储','yali2024.html及txt为开放论文原始网页与文本；moa和ucscab为官方网页快照及文本。yali2025.download-blocked.html是反机器人验证页，不是论文全文。2025正文经浏览工具读取，其结构化知识已整理，但不声称本地已取得该文完整原件。'),
('数据集用途','reviewed_claims.jsonl为经本次资料核对的知识关系及条件；sources.jsonl为原件清单与SHA256；audit.json为数据检查；graph-before.json为变更前备份。未重建FAISS或350题评测集，不能声称问答精度已经提高。'),
('前后数量',json.dumps(audit,ensure_ascii=False,indent=2))]
doc=Document();doc.add_heading('泊头鸭梨知识图谱核对与补充记录',0)
lines=['# 泊头鸭梨知识图谱核对与补充记录','日期：2026-09-25']
for title,body in paragraphs:doc.add_heading(title,1);doc.add_paragraph(body);lines.extend(['','## '+title,body])
doc.add_heading('本次参考资料',1)
for key,url in urls.items():doc.add_paragraph(key+'：'+url);lines.append(key+'：'+url)
doc.save(OUT/'知识图谱核对与补充报告.docx');(OUT/'README.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=True))
